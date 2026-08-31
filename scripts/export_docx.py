import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_font(run, font_name="Times New Roman", east_asia_font="宋体", size_pt=11, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia_font)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)

def render_inline(paragraph, text, base_font="Times New Roman", base_ea="宋体", base_size=11, color_rgb=None):
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            set_font(run, font_name=base_font, east_asia_font=base_ea, size_pt=base_size, bold=True, color_rgb=color_rgb)
        else:
            run = paragraph.add_run(tok)
            set_font(run, font_name=base_font, east_asia_font=base_ea, size_pt=base_size, bold=False, color_rgb=color_rgb)

def render_table(doc, table_lines):
    if len(table_lines) < 2:
        return
    
    header_raw = table_lines[0]
    data_raw = [l for l in table_lines[1:] if not re.match(r"^\|[\s\-:]+\|$", l)]
    
    headers = [c.strip() for c in header_raw.strip("|").split("|")]
    rows_data = [[c.strip() for c in r.strip("|").split("|")] for r in data_raw]
    
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 填充表头
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        clean_h = h_text.replace("**", "")
        run = p.add_run(clean_h)
        set_font(run, font_name="Arial", east_asia_font="黑体", size_pt=9.5, bold=True, color_rgb=RGBColor(24, 43, 73))
        
        shd = parse_xml('<w:shd ' + nsdecls('w') + ' w:fill="EAEFF5"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        
    # 填充数据行
    for row_idx, r_data in enumerate(rows_data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, val in enumerate(r_data):
            if col_idx >= len(row_cells):
                continue
            cell = row_cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            render_inline(p, val, base_size=9.5)
            
    # 设置细边框三线表风格
    tblPr = table._tbl.tblPr
    borders = parse_xml('<w:tblBorders ' + nsdecls('w') + '><w:top w:val="single" w:sz="12" w:space="0" w:color="2C3E50"/><w:bottom w:val="single" w:sz="12" w:space="0" w:color="2C3E50"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(md_path, docx_path):
    doc = Document()
    
    # 设置页面为标准 A4 单栏
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_table = False
    table_lines = []
    
    for line in lines:
        line_str = line.strip()
        if not line_str:
            if in_table:
                render_table(doc, table_lines)
                in_table = False
                table_lines = []
            continue
            
        if line_str.startswith("|") and line_str.endswith("|"):
            in_table = True
            table_lines.append(line_str)
            continue
        elif in_table:
            render_table(doc, table_lines)
            in_table = False
            table_lines = []
            
        # 处理主标题
        if line_str.startswith("# 面向"):
            title_text = line_str.replace("#", "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(18)
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(title_text)
            set_font(run, font_name="Arial", east_asia_font="黑体", size_pt=18, bold=True, color_rgb=RGBColor(24, 43, 73))
            
        # 处理摘要占位
        elif line_str.startswith("**摘要**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(16)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            
            run_lbl = p.add_run("摘要  ")
            set_font(run_lbl, font_name="Times New Roman", east_asia_font="黑体", size_pt=10.5, bold=True)
            
            abs_text = line_str.replace("**摘要**", "").strip()
            run_txt = p.add_run(abs_text)
            set_font(run_txt, font_name="Times New Roman", east_asia_font="楷体", size_pt=10.5, italic=True)
            
        # 处理一级标题
        elif line_str.startswith("# "):
            h1_text = line_str[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h1_text)
            set_font(run, font_name="Arial", east_asia_font="黑体", size_pt=14, bold=True, color_rgb=RGBColor(24, 43, 73))
            
        # 处理二级标题
        elif line_str.startswith("## "):
            h2_text = line_str[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            set_font(run, font_name="Arial", east_asia_font="黑体", size_pt=12, bold=True, color_rgb=RGBColor(44, 62, 80))
            
        # 处理图表占位
        elif line_str.startswith("**[") and line_str.endswith("]**"):
            ph_text = line_str[3:-3].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"【{ph_text}】")
            set_font(run, font_name="Arial", east_asia_font="黑体", size_pt=10.5, bold=True, color_rgb=RGBColor(180, 40, 20))
            
        # 处理图注引用块
        elif line_str.startswith("> "):
            quote_text = line_str[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.line_spacing = 1.2
            
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml('<w:shd ' + nsdecls('w') + ' w:fill="F4F6F9"/>')
            pPr.append(shd)
            
            run = p.add_run(quote_text)
            set_font(run, font_name="Times New Roman", east_asia_font="楷体", size_pt=10.0, italic=False, color_rgb=RGBColor(80, 85, 95))
            
        # 普通正文段落
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.first_line_indent = Inches(0.28) # 约 2 字符首行缩进
            
            render_inline(p, line_str, base_font="Times New Roman", base_ea="宋体", base_size=11)
            
    if in_table:
        render_table(doc, table_lines)
        
    doc.save(docx_path)
    print(f"Docx export complete: {docx_path}")

if __name__ == "__main__":
    md_file = ".scratch/fl-testbed-magazine-draft/prototypes/IEEE-Communications-Magazine中文初稿.md"
    docx_file = ".scratch/fl-testbed-magazine-draft/prototypes/IEEE-Communications-Magazine中文初稿.docx"
    build_docx(md_file, docx_file)
